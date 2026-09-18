from pathlib import Path
from datetime import date

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
CAPTURES = ROOT / "documentacion" / "capturas"
OUTPUT = ROOT / "documentacion" / "Manual_funcional_Monitor_de_Asistencia.docx"

NAVY = "0B253B"
BLUE = "0A6ED1"
PALE_BLUE = "EAF4FC"
PALE_GRAY = "F5F7F9"
GRID = "D9E1E8"
TEXT = "1D2D3E"
MUTED = "526B84"


SCREENS = [
    {
        "title": "Resumen por zona",
        "route": "#/",
        "image": "01-resumen-por-zona.png",
        "purpose": "Concentra la eficiencia operativa, las incidencias y los centros que requieren atención.",
        "filters": "Semana, mes o año; fecha de referencia; horario; estatus.",
        "data": "Asistencia e ineficiencia por zona; indicadores críticos; tendencia de retardos, ausencias y fuera de zona.",
        "actions": "Aplicar filtros; abrir centros de una zona; abrir un centro crítico.",
    },
    {
        "title": "Centros de trabajo",
        "route": "#/centers/{zoneId}",
        "image": "02-centros-de-trabajo.png",
        "purpose": "Compara centros por zona y permite acceder a su operación detallada.",
        "filters": "Zona; semana, mes o año; fecha; búsqueda por nombre o código.",
        "data": "Centros con registros; usuarios; presentes; incidencias; porcentaje de asistencia por centro.",
        "actions": "Abrir centro; alta de centro; activar o desactivar; seleccionar una barra de la gráfica.",
    },
    {
        "title": "Detalle del centro de trabajo",
        "route": "#/center/{centerId}",
        "image": "03-detalle-centro.png",
        "purpose": "Analiza asistencia, puntualidad, incidencias, horarios y registros del centro seleccionado.",
        "filters": "Semana, mes o año; fecha; horario; estatus; búsqueda de usuario.",
        "data": "Asistencia efectiva; puntualidad; incidencias; distribución por horario; tendencia diaria.",
        "actions": "Abrir registros; abrir desempeño del usuario; exportar; ir al alta de usuario.",
    },
    {
        "title": "Registros de asistencia",
        "route": "#/records/{employeeId}",
        "image": "04-registros-asistencia.png",
        "purpose": "Consulta entradas, salidas, comida, horas, distancia, estatus e historial de cada marcación.",
        "filters": "Semana, mes o año; fecha; centro; horario; estatus.",
        "data": "Porcentajes por estatus; tabla paginada; detalle de horario; inicio y fin de comida; distancia.",
        "actions": "Seleccionar registro; ver desempeño; editar; exportar; registrar incidencia.",
    },
    {
        "title": "Turnos y horarios",
        "route": "#/shifts",
        "image": "05-turnos-horarios.png",
        "purpose": "Administra los 21 horarios semanales y revisa su cobertura dentro del periodo elegido.",
        "filters": "Semana, mes o año; fecha; búsqueda; tipo de jornada.",
        "data": "Cobertura del periodo; jornadas evaluadas; ausencias; horarios con menor cobertura; catálogo.",
        "actions": "Filtrar catálogo; seleccionar gráfica; alta de horario con hasta tres reglas semanales.",
    },
    {
        "title": "Administración de usuarios",
        "route": "#/users",
        "image": "06-usuarios.png",
        "purpose": "Gestiona colaboradores, centros, roles, horarios y estatus de activación.",
        "filters": "Búsqueda por nombre, número, centro u horario; horario; estatus activo o inactivo.",
        "data": "Totales de usuarios; centro; horario semanal; rol; estatus; última actualización.",
        "actions": "Alta de usuario; abrir desempeño; abrir mapa; activar o desactivar; exportar.",
    },
    {
        "title": "Desempeño individual",
        "route": "#/user-performance/{employeeId}",
        "image": "07-desempeno-usuario.png",
        "purpose": "Mide la evolución semanal, mensual y anual de un colaborador.",
        "filters": "Vista semanal, mensual o anual; año; mes; semana.",
        "data": "Asistencia efectiva; puntualidad; jornadas completas; cumplimiento de zona; tendencia y detalle.",
        "actions": "Cambiar periodo; revisar registros del periodo; abrir la ubicación del usuario.",
    },
    {
        "title": "Mapa general de ubicaciones",
        "route": "#/map",
        "image": "08-mapa-general.png",
        "purpose": "Presenta las zonas operativas y el estado general de centros y personal.",
        "filters": "Selector de vista geográfica y actualización de ubicaciones.",
        "data": "Zonas; centros; personas visibles; última actualización; leyenda de estados.",
        "actions": "Abrir una zona desde el mapa o la lista; actualizar las ubicaciones.",
    },
    {
        "title": "Mapa de una zona",
        "route": "#/map/{zoneId}",
        "image": "09-mapa-zona.png",
        "purpose": "Ubica centros y personal dentro de la zona seleccionada.",
        "filters": "Zona activa y actualización de ubicaciones.",
        "data": "Centros; técnicos y personal; última señal; estado operativo; precisión disponible.",
        "actions": "Seleccionar centro; seleccionar persona; regresar al mapa general.",
    },
    {
        "title": "Ubicación de un usuario",
        "route": "#/map/{zoneId}/{employeeId}",
        "image": "10-mapa-usuario.png",
        "purpose": "Centra el mapa en un colaborador y muestra la ubicación recibida para su seguimiento.",
        "filters": "Usuario y zona definidos por la navegación.",
        "data": "Nombre; número de empleado; centro; estado; última señal; precisión de la coordenada.",
        "actions": "Volver a la zona; actualizar ubicaciones; regresar al mapa general.",
    },
    {
        "title": "Configuración",
        "route": "#/settings",
        "image": "10-configuracion.png",
        "purpose": "Administra las reglas generales usadas para calcular el estado de asistencia.",
        "filters": "No aplica.",
        "data": "Tolerancia de retardo; entrada anticipada; ausencia automática; frecuencia de actualización.",
        "actions": "Guardar cambios o restablecer los valores de demostración.",
    },
]


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_borders(cell, color=GRID, size="6"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top=80, start=100, bottom=80, end=100):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_run_font(run, name="Arial", size=9, color=TEXT, bold=False):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    run.font.bold = bold


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Página ")
    set_run_font(run, size=8, color=MUTED)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(instr)
    run._r.append(end)


def configure_document(doc):
    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Inches(11)
    section.page_height = Inches(8.5)
    section.top_margin = Inches(0.45)
    section.bottom_margin = Inches(0.45)
    section.left_margin = Inches(0.5)
    section.right_margin = Inches(0.5)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    normal.font.size = Pt(9)
    normal.font.color.rgb = RGBColor.from_string(TEXT)
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.line_spacing = 1.08

    title = styles["Title"]
    title.font.name = "Arial"
    title._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    title._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    title.font.size = Pt(27)
    title.font.bold = True
    title.font.color.rgb = RGBColor(0, 0, 0)
    title_ppr = title._element.get_or_add_pPr()
    title_border = title_ppr.find(qn("w:pBdr"))
    if title_border is not None:
        title_ppr.remove(title_border)

    for style_name, size in (("Heading 1", 16), ("Heading 2", 12)):
        style = styles[style_name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(5)
        style.paragraph_format.space_after = Pt(5)

    footer = section.footer
    add_page_number(footer.paragraphs[0])


def add_screen_summary_table(doc, screen):
    table = doc.add_table(rows=2, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    headers = ("Filtros y alcance", "Datos e indicadores", "Acciones y navegación")
    values = (screen["filters"], screen["data"], screen["actions"])
    for index, header in enumerate(headers):
        table.columns[index].width = Inches(3.2)
        cell = table.cell(0, index)
        cell.width = Inches(3.2)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        set_cell_shading(cell, NAVY)
        set_cell_borders(cell)
        set_cell_margins(cell, top=45, bottom=45)
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_after = Pt(0)
        run = paragraph.add_run(header)
        set_run_font(run, size=8, color="FFFFFF", bold=True)
    set_repeat_table_header(table.rows[0])

    for index, value in enumerate(values):
        cell = table.cell(1, index)
        cell.width = Inches(3.2)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        set_cell_shading(cell, "FFFFFF" if index % 2 == 0 else PALE_GRAY)
        set_cell_borders(cell)
        set_cell_margins(cell, top=55, bottom=55)
        paragraph = cell.paragraphs[0]
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.line_spacing = 1.0
        run = paragraph.add_run(value)
        set_run_font(run, size=7.5, color=TEXT)


def add_screen_page(doc, screen):
    doc.add_page_break()
    heading = doc.add_heading(screen["title"], level=1)
    heading.paragraph_format.keep_with_next = True

    intro = doc.add_paragraph()
    intro.paragraph_format.space_after = Pt(4)
    route = intro.add_run(f"Ruta {screen['route']}  ")
    set_run_font(route, size=8.5, color=BLUE, bold=True)
    purpose = intro.add_run(screen["purpose"])
    set_run_font(purpose, size=8.5, color=MUTED)

    image_path = CAPTURES / screen["image"]
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(4)
    picture = paragraph.add_run().add_picture(str(image_path), width=Inches(9.15))
    picture._inline.docPr.set("descr", f"Pantalla completa de {screen['title']}")

    add_screen_summary_table(doc, screen)


def add_standard_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for index, (header, width) in enumerate(zip(headers, widths)):
        cell = table.cell(0, index)
        cell.width = Inches(width)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        set_cell_shading(cell, NAVY)
        set_cell_borders(cell)
        set_cell_margins(cell)
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_after = Pt(0)
        run = paragraph.add_run(header)
        set_run_font(run, size=8.5, color="FFFFFF", bold=True)
    set_repeat_table_header(table.rows[0])

    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        for col_index, (value, width) in enumerate(zip(values, widths)):
            cell = cells[col_index]
            cell.width = Inches(width)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_shading(cell, "FFFFFF" if row_index % 2 == 0 else PALE_BLUE)
            set_cell_borders(cell)
            set_cell_margins(cell)
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            run = paragraph.add_run(str(value))
            set_run_font(run, size=8, color=TEXT)
    return table


def build_manual():
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    configure_document(doc)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(52)
    title.paragraph_format.space_after = Pt(12)
    title.add_run("Manual funcional del Monitor de Asistencia")

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(24)
    run = subtitle.add_run("Pantallas, funciones, reglas de cálculo y navegación del prototipo SAPUI5")
    set_run_font(run, size=15, color=MUTED)

    opening = doc.add_paragraph()
    opening.alignment = WD_ALIGN_PARAGRAPH.CENTER
    opening.paragraph_format.left_indent = Inches(1.25)
    opening.paragraph_format.right_indent = Inches(1.25)
    opening.paragraph_format.space_after = Pt(18)
    run = opening.add_run(
        "El prototipo cubre el seguimiento por zona, centro y colaborador; la administración de horarios y usuarios; "
        "los registros de asistencia; y la consulta geográfica del personal. Cada sección de este manual incluye la "
        "pantalla completa y resume los filtros, indicadores y acciones disponibles."
    )
    set_run_font(run, size=11, color=TEXT)

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = meta.add_run("Versión documentada 17 de septiembre de 2026")
    set_run_font(run, size=9, color=MUTED, bold=True)

    doc.add_page_break()
    doc.add_heading("Contenido del manual", level=1)
    intro = doc.add_paragraph(
        "Las rutas usan el identificador dinámico de la zona, el centro o el empleado cuando corresponde. "
        "Las pantallas se muestran con datos de demostración al 31 de agosto de 2026."
    )
    intro.paragraph_format.space_after = Pt(8)
    rows = [(str(index), screen["title"], screen["route"], screen["purpose"]) for index, screen in enumerate(SCREENS, 1)]
    add_standard_table(doc, ("Núm.", "Pantalla", "Ruta", "Finalidad"), rows, (0.55, 2.2, 2.15, 5.0))

    for screen in SCREENS:
        add_screen_page(doc, screen)

    doc.add_page_break()
    doc.add_heading("Reglas de cálculo", level=1)
    p = doc.add_paragraph(
        "Los porcentajes visibles muestran su numerador y denominador para que el resultado pueda verificarse desde la interfaz."
    )
    p.paragraph_format.space_after = Pt(8)
    formulas = [
        ("Asistencia efectiva", "Jornadas con entrada ÷ jornadas laborables sin vacaciones ni paternidad × 100", "Centro y desempeño individual"),
        ("Puntualidad", "Entradas sin retardo ÷ asistencias registradas × 100", "Centro y desempeño individual"),
        ("Jornadas completas", "Registros con entrada y salida ÷ asistencias registradas × 100", "Desempeño individual"),
        ("Cumplimiento de zona", "Marcaciones con distancia menor o igual a 2 m ÷ marcaciones con GPS × 100", "Desempeño individual"),
        ("Ineficiencia", "Retardos y ausencias ÷ total de registros × 100", "Resumen por zona"),
        ("Cobertura de horario", "Registros presentes ÷ jornadas evaluadas × 100", "Turnos y horarios"),
        ("Porcentaje por estatus", "Cantidad del estatus ÷ total filtrado × 100", "Registros de asistencia"),
    ]
    add_standard_table(doc, ("Indicador", "Fórmula", "Pantallas"), formulas, (2.0, 5.0, 2.9))

    doc.add_heading("Reglas operativas", level=2)
    rules = [
        "La semana se calcula de lunes a domingo a partir de la fecha de referencia.",
        "La vista mensual usa el mes de la fecha seleccionada y la anual usa su año.",
        "Una distancia superior a 2 metros se clasifica como Fuera de zona, salvo viaje de trabajo, vacaciones, paternidad y home office.",
        "Inicio de comida representa la salida temporal; fin de comida representa el regreso a la jornada.",
        "Los horarios que terminan a una hora menor o igual a la inicial se interpretan como jornadas que concluyen al día siguiente.",
    ]
    for rule in rules:
        paragraph = doc.add_paragraph(style=None)
        paragraph.style = doc.styles["Normal"]
        paragraph.paragraph_format.left_indent = Inches(0.2)
        paragraph.paragraph_format.first_line_indent = Inches(-0.15)
        paragraph.add_run("• ")
        paragraph.add_run(rule)

    doc.add_page_break()
    doc.add_heading("Estados y datos de demostración", level=1)
    statuses = [
        ("Presente", "Entrada y jornada registradas sin incidencia principal."),
        ("Retardo", "Entrada posterior a la tolerancia configurada."),
        ("Ausente", "No existe entrada válida para la jornada."),
        ("Incompleto", "Falta una marcación necesaria para cerrar la jornada."),
        ("Viaje de trabajo", "Actividad autorizada fuera del centro habitual."),
        ("Vacaciones", "Permiso autorizado que se excluye del denominador de asistencia efectiva."),
        ("Paternidad", "Permiso autorizado que se excluye del denominador de asistencia efectiva."),
        ("Home office", "Jornada autorizada fuera del centro y considerada asistencia."),
        ("Fuera de zona", "La distancia registrada supera los 2 metros permitidos."),
    ]
    add_standard_table(doc, ("Estatus", "Definición funcional"), statuses, (2.3, 7.6))

    doc.add_heading("Alcance actual del prototipo", level=2)
    paragraphs = [
        "El modelo local contiene cuatro zonas operativas, 16 centros, 18 usuarios, 21 horarios semanales, registros recientes e historial anual, además de ubicaciones de prueba para el mapa.",
        "Las altas y los cambios modifican el modelo local. Exportar, nueva incidencia y editar registro están preparados para conectarse a los servicios SAP correspondientes.",
        "El mapa utiliza SAPUI5 Visual Business con mosaicos de OpenStreetMap. En producción, las coordenadas y sus marcas de tiempo deberán recibirse desde SAP y respetar los permisos de ubicación definidos por la organización.",
    ]
    for text in paragraphs:
        doc.add_paragraph(text)

    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_manual()

